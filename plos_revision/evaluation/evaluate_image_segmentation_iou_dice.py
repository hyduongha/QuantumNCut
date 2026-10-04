# -*- coding: utf-8 -*-

"""

Evaluate IoU and Dice for ground truth and the selected prediction methods:

    Ncut, Quantum

Features:

- Stream .seg files row by row without allocating an H x W mask in memory.

- The end_col coordinate is inclusive.

- Label 0 is a valid segmentation class and is included in the class average.

- Each ground-truth/prediction pair is scanned once for both IoU and Dice.

- Each dataset occupies one row in the Excel output.

- Each run creates a new workbook and overwrites the previous result file.

"""

import os

import time

from collections import defaultdict

from pathlib import Path

import numpy as np

from openpyxl import Workbook

from openpyxl.styles import Alignment, Font

from openpyxl.utils import get_column_letter

from scipy.optimize import linear_sum_assignment

# ============================================================

# 1. CONFIGURATION

# ============================================================

# Directory containing the merged ground-truth .seg files

GT_ROOT = Path(

    r"F:\Nhu Y_khong xoa_2\groundtruth_masks_MaSS13K"

)

# Directory containing the prediction datasets

PREDICTION_ROOT = Path(

    r"D:\NhuY_khongxoa_4_tmp\Quantum_QPE_QL_IQPE_1_40"

)

# Evaluate only the datasets listed here

DATASET_FOLDER_NAMES = (

    "05823_b06_20230210_tianjin_id2965_2928x3904_K4",

    "05913_b06_20230210_tianjin_id2968_3648x2736_K4",

    "06115_b06_20230210_tianjin_id2974_2736x3648_K3",

    "06875_b06_20230210_tianjin_id3000_2448x3264_K6",

    "06888_b06_20230210_tianjin_id3000_3000x4000_K4",

    "06934_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06935_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06936_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06944_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06956_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06957_b06_20230210_tianjin_id3002_2736x3648_K4",

    "08895_b07_20230217_tianjin_id4493_3648x2736_K4",

    "09007_b07_20230217_tianjin_id4497_3648x2736_K4",

    "09653_b07_20230217_chengdu_id4518_3072x4096_K4",

    "12550_b07_20230217_chengdu_id4537_3648x2736_K4",

    "12791_b07_20230217_chengdu_id4545_3000x4000_K4",

    "12914_b07_20230217_chengdu_id4549_3000x4000_K4",

    "13057_b08_20230224_yinchuan_id3080_3000x4000_K4",

    "13128_b08_20230224_yinchuan_id3082_3000x4000_K4",

    "13421_b08_20230224_yinchuan_id3092_3000x4000_K4",

    "14034_b08_20230224_chengdu_id3113_3000x4000_K4",

    "14064_b08_20230224_chengdu_id3114_3000x4000_K4",

    "14068_b08_20230224_chengdu_id3114_3000x4000_K4",

    "14133_b08_20230224_chengdu_id3117_3000x4000_K4",

    "15311_b08_20230224_yinchuan_id3122_3000x4000_K4",

    "15376_b08_20230224_yinchuan_id3124_3000x4000_K4",

    "15531_b08_20230224_yinchuan_id3129_2736x3648_K4",

    "15537_b08_20230224_yinchuan_id3129_3000x4000_K4",

    "15556_b08_20230224_yinchuan_id3130_3000x4000_K4",

    "15585_b08_20230224_yinchuan_id3131_3000x4000_K4",

    "15659_b08_20230224_yinchuan_id3133_2736x3648_K4",

    "15765_b08_20230224_yinchuan_id3137_2736x3648_K5",

    "15922_b08_20230224_yinchuan_id3142_3000x4000_K4",

    "15976_b08_20230224_yinchuan_id3144_2736x3648_K4",

    "19364_b09_20230303_yinchuan_id5653_3024x4032_K3",

    "24778_b09_20230303_yinchuan_id5916_3000x4000_K3",

    "24992_b09_20230303_yinchuan_id5923_2736x3648_K5",

    "26075_b09_20230303_yinchuan_id5959_2736x3648_K4",

    "29055_b10_20230310_chengdu_id6061_2928x3904_K4",

    "56803_b26_20230703_tianjin_id8154_4096x3072_K3",

)

# Subdirectory containing merged prediction .seg files for each dataset

MERGED_FOLDER_NAME = "seg_merge"

# Methods to evaluate (uncomment additional methods as needed)

METHOD_NAMES = (

#    "QL",

    "Ncut",

#    "QPE",

#    "IQPE",

    "Quantum",

)

# Prediction filename pattern:

#   <dataset>_QL.seg

#   <dataset>_Ncut.seg

#   <dataset>_QPE.seg

#   <dataset>_IQPE.seg

#   <dataset>_Quantum.seg

PREDICTION_FILE_TEMPLATE = "{dataset}_{method}.seg"

# Set to False when prediction labels already match ground-truth labels

USE_MAPPING = False

# Progress logging intervals; set either value to 0 to disable its messages

LOG_EVERY_ROWS = 500

LOG_EVERY_LINES = 2_000_000

# The Excel output file is overwritten on each run

OUTPUT_EXCEL = PREDICTION_ROOT / "evaluation_IoU_Dice_Mcut_Quantum.xlsx"

SHEET_NAME = "Results"

EXCEL_COLUMNS = (

    "dataset",

    "IoU_QL",

    "IoU_Ncut",

    "IoU_QPE",

    "IoU_IQPE",

    "IoU_Quantum",

    "Dice_QL",

    "Dice_Ncut",

    "Dice_QPE",

    "Dice_IQPE",

    "Dice_Quantum",

)

# ============================================================

# 2. READ SEG HEADERS AND ROW DATA

# ============================================================

def read_seg_dimensions(seg_path: Path):

    """Read width and height from a .seg file header."""

    width = None

    height = None

    with open(seg_path, "r", encoding="utf-8", errors="replace") as file:

        for line in file:

            text = line.strip()

            lower_text = text.lower()

            if lower_text.startswith("width "):

                width = int(text.split()[1])

            elif lower_text.startswith("height "):

                height = int(text.split()[1])

            elif lower_text == "data":

                break

    if width is None or height is None:

        raise ValueError(f"SEG file is missing width or height:\n{seg_path}")

    return width, height

def iter_seg_rows(seg_path: Path, log_every_lines: int = 2_000_000):

    """

    Yield:

        row_index, [(label, start_col, end_col), ...]

    end_col is inclusive.

    Rows must be sorted in ascending order, and runs within each row

    must be sorted by start_col in ascending order.

    """

    width = None

    in_data = False

    current_row = None

    runs = []

    data_line_count = 0

    with open(seg_path, "r", encoding="utf-8", errors="replace") as file:

        for line in file:

            if not in_data:

                text = line.strip()

                lower_text = text.lower()

                if lower_text == "data":

                    in_data = True

                    continue

                if lower_text.startswith("width "):

                    width = int(text.split()[1])

                continue

            parts = line.split()

            if len(parts) != 4:

                continue

            try:

                label = int(parts[0])

                row = int(parts[1])

                start_col = int(parts[2])

                end_col = int(parts[3])

            except ValueError:

                continue

            data_line_count += 1

            if log_every_lines and data_line_count % log_every_lines == 0:

                print(

                    f"[{seg_path.name}] read "

                    f"{data_line_count:,} data lines..."

                )

            # Clamp column indices to the image width

            if width is not None:

                if start_col < 0:

                    start_col = 0

                if end_col >= width:

                    end_col = width - 1

            if start_col > end_col:

                continue

            if current_row is None:

                current_row = row

            if row != current_row:

                yield current_row, runs

                current_row = row

                runs = []

            runs.append((label, start_col, end_col))

    if in_data and current_row is not None:

        yield current_row, runs

# ============================================================

# 3. ACCUMULATE OVERLAP

# ============================================================

def accumulate_overlap(

    seg1_path: Path,

    seg2_path: Path,

    log_every_rows: int = 500,

    log_every_lines: int = 2_000_000,

):

    """

    Returns:

        overlap[(label_gt, label_pred)] = number of overlapping pixels

        fg1, fg2 = number of labeled pixels in each file (including label 0)

        both_fg = number of pixels labeled in both files

    """

    iterator1 = iter_seg_rows(seg1_path, log_every_lines)

    iterator2 = iter_seg_rows(seg2_path, log_every_lines)

    overlap = defaultdict(int)

    foreground1 = 0

    foreground2 = 0

    both_foreground = 0

    row1, runs1 = next(iterator1, (None, None))

    row2, runs2 = next(iterator2, (None, None))

    rows_done = 0

    while row1 is not None and row2 is not None:

        if row1 < row2:

            for label1, start1, end1 in runs1:

                if label1 >= 0:

                    foreground1 += end1 - start1 + 1

            row1, runs1 = next(iterator1, (None, None))

            continue

        if row2 < row1:

            for label2, start2, end2 in runs2:

                if label2 >= 0:

                    foreground2 += end2 - start2 + 1

            row2, runs2 = next(iterator2, (None, None))

            continue

        # row1 == row2

        rows_done += 1

        if log_every_rows and rows_done % log_every_rows == 0:

            print(

                f"Processed {rows_done:,} rows "

                f"(current row = {row1})..."

            )

        for label1, start1, end1 in runs1:

            if label1 >= 0:

                foreground1 += end1 - start1 + 1

        for label2, start2, end2 in runs2:

            if label2 >= 0:

                foreground2 += end2 - start2 + 1

        # Use two pointers to find overlap between runs on the same row

        index1 = 0

        index2 = 0

        while index1 < len(runs1) and index2 < len(runs2):

            label1, start1, end1 = runs1[index1]

            label2, start2, end2 = runs2[index2]

            overlap_start = max(start1, start2)

            overlap_end = min(end1, end2)

            if overlap_start <= overlap_end:

                pixel_count = overlap_end - overlap_start + 1

                overlap[(label1, label2)] += pixel_count

                if label1 >= 0 and label2 >= 0:

                    both_foreground += pixel_count

            if end1 < end2:

                index1 += 1

            else:

                index2 += 1

        row1, runs1 = next(iterator1, (None, None))

        row2, runs2 = next(iterator2, (None, None))

    # Remaining ground-truth rows

    while row1 is not None:

        for label1, start1, end1 in runs1:

            if label1 >= 0:

                foreground1 += end1 - start1 + 1

        row1, runs1 = next(iterator1, (None, None))

    # Remaining prediction rows

    while row2 is not None:

        for label2, start2, end2 in runs2:

            if label2 >= 0:

                foreground2 += end2 - start2 + 1

        row2, runs2 = next(iterator2, (None, None))

    return overlap, foreground1, foreground2, both_foreground

# ============================================================

# 4. HUNGARIAN LABEL MATCHING AND METRICS

# ============================================================

def hungarian_mapping_from_overlap(overlap_dictionary):

    """Map prediction labels to ground-truth labels."""

    labels1 = sorted({label1 for label1, _ in overlap_dictionary})

    labels2 = sorted({label2 for _, label2 in overlap_dictionary})

    if not labels1 or not labels2:

        return {}

    index1 = {label: index for index, label in enumerate(labels1)}

    index2 = {label: index for index, label in enumerate(labels2)}

    confusion_matrix = np.zeros(

        (len(labels1), len(labels2)),

        dtype=np.int64,

    )

    for (label1, label2), pixel_count in overlap_dictionary.items():

        confusion_matrix[index1[label1], index2[label2]] = pixel_count

    row_indices, column_indices = linear_sum_assignment(-confusion_matrix)

    return {

        labels2[column_index]: labels1[row_index]

        for row_index, column_index in zip(row_indices, column_indices)

    }

def calculate_iou_and_dice(

    overlap,

    foreground1,

    foreground2,

    both_foreground,

    use_mapping=False,

):

    """Compute mIoU and macro-Dice (%) across all labels, including label 0.

    Labels present in ground truth or prediction are included in the average.

    When USE_MAPPING=True, match prediction labels to ground truth by overlap;

    unmatched prediction labels form separate classes to account for false positives.

    """

    mapping = hungarian_mapping_from_overlap(overlap) if use_mapping else None

    gt_labels = {label_gt for label_gt, _ in overlap}

    pred_labels = {label_pred for _, label_pred in overlap}

    # Include unmatched ground-truth and prediction labels in the scores.

    if use_mapping:

        pred_to_class = {

            label: mapping.get(label, ("unmatched_pred", label))

            for label in pred_labels

        }

    else:

        pred_to_class = {label: label for label in pred_labels}

    gt_size = defaultdict(int)

    pred_size = defaultdict(int)

    intersection = defaultdict(int)

    for (label_gt, label_pred), pixel_count in overlap.items():

        pred_class = pred_to_class[label_pred]

        gt_size[label_gt] += pixel_count

        pred_size[pred_class] += pixel_count

        if label_gt == pred_class:

            intersection[label_gt] += pixel_count

    classes = gt_labels | set(pred_to_class.values())

    if not classes:

        raise ValueError("No overlapping labeled pixels were found for the two SEG files.")

    iou_scores = []

    dice_scores = []

    for label in classes:

        tp = intersection[label]

        gt_count = gt_size[label]

        pred_count = pred_size[label]

        union = gt_count + pred_count - tp

        iou_scores.append(tp / union if union else 0.0)

        dice_scores.append(2 * tp / (gt_count + pred_count)

                           if gt_count + pred_count else 0.0)

    return (sum(iou_scores) / len(classes) * 100.0,

            sum(dice_scores) / len(classes) * 100.0,

            mapping)

# ============================================================

def find_ground_truth_seg(dataset_name: str):

    """Find the exact <dataset>.seg file under GT_ROOT."""

    direct_path = GT_ROOT / f"{dataset_name}.seg"

    if direct_path.is_file():

        return direct_path

    nested_path = GT_ROOT / dataset_name / f"{dataset_name}.seg"

    if nested_path.is_file():

        return nested_path

    matches = list(GT_ROOT.rglob(f"{dataset_name}.seg"))

    if len(matches) == 1:

        return matches[0]

    if len(matches) > 1:

        print("Multiple ground-truth files share this name; no file was selected:")

        for match in matches:

            print(f"   {match}")

    return None

def build_prediction_path(dataset_name: str, method_name: str):

    """Build the merged prediction file path for one method."""

    file_name = PREDICTION_FILE_TEMPLATE.format(

        dataset=dataset_name,

        method=method_name,

    )

    return (

        PREDICTION_ROOT

        / dataset_name

        / MERGED_FOLDER_NAME

        / file_name

    )

def evaluate_one_pair(ground_truth_path: Path, prediction_path: Path):

    """Evaluate IoU and Dice for one ground-truth/prediction pair."""

    gt_width, gt_height = read_seg_dimensions(ground_truth_path)

    pred_width, pred_height = read_seg_dimensions(prediction_path)

    if (gt_width, gt_height) != (pred_width, pred_height):

        raise ValueError(

            "Ground-truth and prediction dimensions do not match:\n"

            f"GT   : {gt_width}x{gt_height}\n"

            f"Pred : {pred_width}x{pred_height}"

        )

    print(

        f"GT size   : "

        f"{os.path.getsize(ground_truth_path) / 1024 / 1024:.2f} MB"

    )

    print(

        f"Pred size : "

        f"{os.path.getsize(prediction_path) / 1024 / 1024:.2f} MB"

    )

    start_time = time.time()

    overlap, foreground1, foreground2, both_foreground = accumulate_overlap(

        ground_truth_path,

        prediction_path,

        log_every_rows=LOG_EVERY_ROWS,

        log_every_lines=LOG_EVERY_LINES,

    )

    iou, dice, mapping = calculate_iou_and_dice(

        overlap,

        foreground1,

        foreground2,

        both_foreground,

        use_mapping=USE_MAPPING,

    )

    return {

        "iou": iou,

        "dice": dice,

        "mapping": mapping,

        "seconds": time.time() - start_time,

    }

def empty_result_row(dataset_name: str):

    """Create a result row with empty metric values."""

    result = {"dataset": dataset_name}

    for method_name in METHOD_NAMES:

        result[f"IoU_{method_name}"] = None

    for method_name in METHOD_NAMES:

        result[f"Dice_{method_name}"] = None

    return result

def evaluate_one_dataset(dataset_name: str):

    """Evaluate the selected methods for one dataset."""

    result_row = empty_result_row(dataset_name)

    print("\n" + "=" * 100)

    print(f"EVALUATING DATASET: {dataset_name}")

    print("=" * 100)

    dataset_folder = PREDICTION_ROOT / dataset_name

    merged_folder = dataset_folder / MERGED_FOLDER_NAME

    if not dataset_folder.is_dir():

        print(f"Dataset directory not found:\n   {dataset_folder}")

        return result_row

    if not merged_folder.is_dir():

        print(f"Directory not found:\n   {merged_folder}")

        return result_row

    ground_truth_path = find_ground_truth_seg(dataset_name)

    if ground_truth_path is None:

        print(f"Ground-truth file not found: {dataset_name}.seg")

        return result_row

    print(f"Ground Truth : {ground_truth_path}")

    print(f"Prediction   : {merged_folder}")

    for method_name in METHOD_NAMES:

        prediction_path = build_prediction_path(dataset_name, method_name)

        print("\n" + "-" * 100)

        print(f"METHOD: {method_name}")

        print(f"GT   : {ground_truth_path}")

        print(f"Pred : {prediction_path}")

        print("-" * 100)

        if not prediction_path.is_file():

            print("Prediction file not found. IoU and Dice cells will remain empty.")

            continue

        try:

            evaluation = evaluate_one_pair(

                ground_truth_path,

                prediction_path,

            )

            result_row[f"IoU_{method_name}"] = evaluation["iou"]

            result_row[f"Dice_{method_name}"] = evaluation["dice"]

            print(f"IoU_{method_name}  : {evaluation['iou']:.6f}%")

            print(f"Dice_{method_name} : {evaluation['dice']:.6f}%")

            print(f"   Elapsed time          : {evaluation['seconds']:.2f} seconds")

            if USE_MAPPING:

                print(f"   Mapping            : {evaluation['mapping']}")

            else:

                print("   Mapping            : disabled")

        except Exception as error:

            print(f"Could not evaluate {method_name}:\n   {error}")

    return result_row

# ============================================================

# 6. WRITE EXCEL RESULTS

# ============================================================

def write_results_to_excel(result_rows, excel_path: Path):

    """

    Always create a new workbook.

    If the Excel file exists, workbook.save() overwrites it.

    """

    excel_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = SHEET_NAME

    # Header

    worksheet.append(list(EXCEL_COLUMNS))

    for cell in worksheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows

    for result_row in result_rows:

        worksheet.append([

            result_row.get(column_name)

            for column_name in EXCEL_COLUMNS

        ])

        current_row = worksheet.max_row

        worksheet.cell(current_row, 1).alignment = Alignment(

            horizontal="left",

            vertical="center",

        )

        for column_index in range(2, len(EXCEL_COLUMNS) + 1):

            cell = worksheet.cell(current_row, column_index)

            if cell.value is not None:

                cell.number_format = "0.000000"

            cell.alignment = Alignment(

                horizontal="center",

                vertical="center",

            )

    # General formatting

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (

        f"A1:{get_column_letter(len(EXCEL_COLUMNS))}{worksheet.max_row}"

    )

    worksheet.column_dimensions["A"].width = 62

    for column_index in range(2, len(EXCEL_COLUMNS) + 1):

        worksheet.column_dimensions[get_column_letter(column_index)].width = 16

    try:

        workbook.save(excel_path)

    except PermissionError as error:

        raise PermissionError(

            "Cannot overwrite the Excel file. It may be open in another program.\n"

            f"Close the file and run the script again:\n{excel_path}"

        ) from error

    print("\nCreated and saved the Excel file (overwriting any existing file):")

    print(f"   {excel_path}")

# ============================================================

# 7. PRINT SUMMARY

# ============================================================

def format_metric(value):

    return "-" if value is None else f"{value:.4f}"

def print_summary(result_rows):

    print("\n" + "=" * 178)

    print("IOU AND DICE SUMMARY FOR SELECTED METHODS")

    print("=" * 178)

    header = f"{'Dataset':<50}"

    for method_name in METHOD_NAMES:

        header += f"{'IoU_' + method_name:>12}"

    for method_name in METHOD_NAMES:

        header += f"{'Dice_' + method_name:>12}"

    print(header)

    print("-" * 178)

    for result in result_rows:

        dataset_display = result["dataset"]

        if len(dataset_display) > 47:

            dataset_display = dataset_display[:44] + "..."

        line = f"{dataset_display:<50}"

        for method_name in METHOD_NAMES:

            line += f"{format_metric(result[f'IoU_{method_name}']):>12}"

        for method_name in METHOD_NAMES:

            line += f"{format_metric(result[f'Dice_{method_name}']):>12}"

        print(line)

# ============================================================

# 8. MAIN

# ============================================================

def main():

    if not GT_ROOT.is_dir():

        raise FileNotFoundError(f"GT_ROOT directory not found:\n{GT_ROOT}")

    if not PREDICTION_ROOT.is_dir():

        raise FileNotFoundError(

            f"PREDICTION_ROOT directory not found:\n{PREDICTION_ROOT}"

        )

    selected_dataset_names = list(dict.fromkeys(

        dataset_name.strip()

        for dataset_name in DATASET_FOLDER_NAMES

        if dataset_name.strip()

    ))

    if not selected_dataset_names:

        print("No datasets were specified in DATASET_FOLDER_NAMES.")

        return

    print("\n" + "=" * 100)

    print("DATASETS TO EVALUATE")

    print("=" * 100)

    for index, dataset_name in enumerate(selected_dataset_names, start=1):

        print(f"{index}. {dataset_name}")

    all_result_rows = []

    total_start_time = time.time()

    for dataset_name in selected_dataset_names:

        all_result_rows.append(evaluate_one_dataset(dataset_name))

    total_seconds = time.time() - total_start_time

    # Create a new Excel file and overwrite the previous one

    write_results_to_excel(all_result_rows, OUTPUT_EXCEL)

    print_summary(all_result_rows)

    print("\n" + "=" * 100)

    print("EVALUATION COMPLETE")

    print("=" * 100)

    print(f"Datasets evaluated : {len(all_result_rows)}")

    print(f"Methods evaluated         : {len(METHOD_NAMES)}")

    print(f"Total elapsed time         : {total_seconds:.2f} seconds")

    print(f"File Excel             : {OUTPUT_EXCEL}")

    print("The next run will overwrite this Excel file.")

if __name__ == "__main__":

    main()
