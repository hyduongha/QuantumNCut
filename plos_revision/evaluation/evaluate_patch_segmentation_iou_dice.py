# -*- coding: utf-8 -*-

"""Compute mIoU and macro-Dice for each .seg patch, then average across patches for each image.

Each prediction patch <dataset>_<patch_id>_<suffix>.seg is paired with

<dataset>_<patch_id>_<other_suffix>.seg in the ground truth. The final suffix is ignored during pairing.

Label 0 is a valid class. All classes present in either ground truth or prediction are included.

"""

from collections import defaultdict

from pathlib import Path

import re

import time

import numpy as np

from openpyxl import Workbook

from openpyxl.styles import Alignment, Font

from openpyxl.utils import get_column_letter

from scipy.optimize import linear_sum_assignment

# ==================== CONFIGURATION ====================

GT_ROOT = Path(r"F:\Nhu Y_khong xoa_2\8x16_MaSS13K\split_masks")

PREDICTION_ROOT = Path(r"D:\NhuY_khongxoa_4_tmp\Quantum_QPE_QL_IQPE_1_40")

PREDICTION_SUBFOLDER = Path("seg_remap")

METHOD_NAMES = ("Ncut", "Quantum")  # Evaluate the two currently enabled methods

DATASET_FOLDER_NAMES = (

    "05823_b06_20230210_tianjin_id2965_2928x3904_K4",

    "05913_b06_20230210_tianjin_id2968_3648x2736_K4",

    "06115_b06_20230210_tianjin_id2974_2736x3648_K3",

    "06875_b06_20230210_tianjin_id3000_2448x3264_K6",

    "06888_b06_20230210_tianjin_id3000_3000x4000_K4",

    "06934_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06935_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06936_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06944_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06956_b06_20230210_tianjin_id3002_2736x3648_K4",

    "06957_b06_20230210_tianjin_id3002_2736x3648_K4",

    "08895_b07_20230217_tianjin_id4493_3648x2736_K4",

    "09007_b07_20230217_tianjin_id4497_3648x2736_K4",

    "09653_b07_20230217_chengdu_id4518_3072x4096_K4",

    "12550_b07_20230217_chengdu_id4537_3648x2736_K4",

    "12791_b07_20230217_chengdu_id4545_3000x4000_K4",

    "12914_b07_20230217_chengdu_id4549_3000x4000_K4",

    "13057_b08_20230224_yinchuan_id3080_3000x4000_K4",

    "13128_b08_20230224_yinchuan_id3082_3000x4000_K4",

    "13421_b08_20230224_yinchuan_id3092_3000x4000_K4",

    "14034_b08_20230224_chengdu_id3113_3000x4000_K4",

    "14064_b08_20230224_chengdu_id3114_3000x4000_K4",

    "14068_b08_20230224_chengdu_id3114_3000x4000_K4",

    "14133_b08_20230224_chengdu_id3117_3000x4000_K4",

    "15311_b08_20230224_yinchuan_id3122_3000x4000_K4",

    "15376_b08_20230224_yinchuan_id3124_3000x4000_K4",

    "15531_b08_20230224_yinchuan_id3129_2736x3648_K4",

    "15537_b08_20230224_yinchuan_id3129_3000x4000_K4",

    "15556_b08_20230224_yinchuan_id3130_3000x4000_K4",

    "15585_b08_20230224_yinchuan_id3131_3000x4000_K4",

    "15659_b08_20230224_yinchuan_id3133_2736x3648_K4",

    "15765_b08_20230224_yinchuan_id3137_2736x3648_K5",

    "15922_b08_20230224_yinchuan_id3142_3000x4000_K4",

    "15976_b08_20230224_yinchuan_id3144_2736x3648_K4",

    "19364_b09_20230303_yinchuan_id5653_3024x4032_K3",

    "24778_b09_20230303_yinchuan_id5916_3000x4000_K3",

    "24992_b09_20230303_yinchuan_id5923_2736x3648_K5",

    "26075_b09_20230303_yinchuan_id5959_2736x3648_K4",

    "29055_b10_20230310_chengdu_id6061_2928x3904_K4",

    "56803_b26_20230703_tianjin_id8154_4096x3072_K3",

)  # Selected image datasets

USE_MAPPING = False  # Set True if prediction labels have not yet been aligned to ground truth

OUTPUT_EXCEL = PREDICTION_ROOT / "evaluation_patch_IoU_Dice_Ncut_Quantum.xlsx"

PRINT_EVERY_PATCHES = 500

# ==================== READ SEG FILES ====================

def read_dimensions(path):

    width = height = None

    with path.open("r", encoding="utf-8", errors="replace") as file:

        for raw in file:

            parts = raw.strip().split()

            if len(parts) == 2 and parts[0].lower() == "width":

                width = int(parts[1])

            elif len(parts) == 2 and parts[0].lower() == "height":

                height = int(parts[1])

            elif parts and parts[0].lower() == "data":

                break

    if not width or not height:

        raise ValueError(f"Missing width/height: {path}")

    return width, height

def iter_rows(path, width, height):

    """Each row contains (label, start, end) runs; end is inclusive."""

    in_data = False

    current_row = None

    runs = []

    with path.open("r", encoding="utf-8", errors="replace") as file:

        for lineno, raw in enumerate(file, 1):

            parts = raw.split()

            if not in_data:

                if parts and parts[0].lower() == "data":

                    in_data = True

                continue

            if not parts:

                continue

            if len(parts) != 4:

                raise ValueError(f"Data line does not contain four columns: {path}:{lineno}")

            try:

                label, row, start, end = map(int, parts)

            except ValueError as exc:

                raise ValueError(f"Data fields must be integers: {path}:{lineno}") from exc

            if not (0 <= row < height and 0 <= start <= end < width):

                raise ValueError(f"Run extends beyond patch dimensions: {path}:{lineno}")

            if current_row is None:

                current_row = row

            if row != current_row:

                if row <= current_row:

                    raise ValueError(f"SEG rows are not strictly increasing: {path}:{lineno}")

                if not runs or runs[0][1] != 0 or runs[-1][2] != width - 1:

                    raise ValueError(f"Row {current_row} does not cover the full width: {path}")

                yield current_row, runs

                current_row, runs = row, []

            if runs and start != runs[-1][2] + 1:

                raise ValueError(f"Gap or overlap between runs in row {row}: {path}:{lineno}")

            runs.append((label, start, end))

    if current_row is None:

        raise ValueError(f"SEG file has no data rows: {path}")

    if runs[0][1] != 0 or runs[-1][2] != width - 1:

        raise ValueError(f"Row {current_row} does not cover the full width: {path}")

    yield current_row, runs

def accumulate_overlap(gt_path, pred_path):

    width, height = read_dimensions(gt_path)

    if (width, height) != read_dimensions(pred_path):

        raise ValueError("Ground truth and prediction have different patch dimensions")

    gt_iter = iter_rows(gt_path, width, height)

    pred_iter = iter_rows(pred_path, width, height)

    overlap = defaultdict(int)

    for expected_row in range(height):

        try:

            gt_row, gt_runs = next(gt_iter)

            pred_row, pred_runs = next(pred_iter)

        except StopIteration as exc:

            raise ValueError("SEG file has fewer rows than specified in the header") from exc

        if gt_row != expected_row or pred_row != expected_row:

            raise ValueError(f"SEG file is missing row {expected_row}")

        i = j = 0

        while i < len(gt_runs) and j < len(pred_runs):

            gt_label, ga, gb = gt_runs[i]

            pred_label, pa, pb = pred_runs[j]

            count = min(gb, pb) - max(ga, pa) + 1

            if count > 0:

                overlap[(gt_label, pred_label)] += count

            if gb <= pb:

                i += 1

            if pb <= gb:

                j += 1

    if next(gt_iter, None) is not None or next(pred_iter, None) is not None:

        raise ValueError("SEG file has more rows than specified in the header")

    if sum(overlap.values()) != width * height:

        raise ValueError("Overlapping pixel count does not equal width x height")

    return overlap

# ==================== PER-CLASS METRICS ====================

def calculate_macro_metrics(overlap, use_mapping=False):

    gt_labels = {g for g, _ in overlap}

    pred_labels = {p for _, p in overlap}

    mapping = None

    if use_mapping:

        gt_order, pred_order = sorted(gt_labels), sorted(pred_labels)

        mat = np.zeros((len(gt_order), len(pred_order)), dtype=np.int64)

        gi = {g: i for i, g in enumerate(gt_order)}

        pi = {p: i for i, p in enumerate(pred_order)}

        for (g, p), count in overlap.items():

            mat[gi[g], pi[p]] = count

        rows, cols = linear_sum_assignment(-mat)

        mapping = {pred_order[c]: gt_order[r] for r, c in zip(rows, cols)}

        # An unmatched prediction label remains a separate class with a score of zero.

        pred_class = {p: mapping.get(p, ("unmatched", p)) for p in pred_labels}

    else:

        pred_class = {p: p for p in pred_labels}

    gt_count, pred_count, tp = defaultdict(int), defaultdict(int), defaultdict(int)

    for (g, p), count in overlap.items():

        c = pred_class[p]

        gt_count[g] += count

        pred_count[c] += count

        if g == c:

            tp[g] += count

    classes = gt_labels | set(pred_class.values())

    per_class = {}

    for label in classes:

        intersect = tp[label]

        gt_n, pred_n = gt_count[label], pred_count[label]

        per_class[label] = (

            intersect / (gt_n + pred_n - intersect),

            2 * intersect / (gt_n + pred_n),

        )

    miou = sum(v[0] for v in per_class.values()) / len(per_class) * 100

    macro_dice = sum(v[1] for v in per_class.values()) / len(per_class) * 100

    return miou, macro_dice, mapping

# ==================== MATCH PATCHES AND EXPORT EXCEL ====================

def patch_index(path, dataset):

    """Extract the patch index; the suffix after it (e.g., _1/_2) may differ."""

    pattern = rf"^{re.escape(dataset)}_(\d+)_([^_]+)\.seg$"

    match = re.fullmatch(pattern, path.name, flags=re.IGNORECASE)

    return match.group(1) if match else None

def index_folder(folder, dataset):

    result = {}

    for path in sorted(folder.glob("*.seg")):

        index = patch_index(path, dataset)

        if index is None:

            continue

        # Normalize 00001 and 1 to the same index; retain filenames for the Excel audit trail.

        key = int(index)

        if key in result:

            raise ValueError(f"Duplicate patch index {key}: {result[key]} and {path}")

        result[key] = path

    return result

def dataset_names():

    if DATASET_FOLDER_NAMES:

        return list(dict.fromkeys(DATASET_FOLDER_NAMES))

    return sorted(p.name for p in PREDICTION_ROOT.iterdir() if p.is_dir())

def make_sheet(workbook, title, headings):

    sheet = workbook.create_sheet(title)

    sheet.append(headings)

    for cell in sheet[1]:

        cell.font = Font(bold=True)

        cell.alignment = Alignment(horizontal="center")

    sheet.freeze_panes = "A2"

    return sheet

def main():

    if not GT_ROOT.is_dir() or not PREDICTION_ROOT.is_dir():

        raise FileNotFoundError(f"Check GT_ROOT={GT_ROOT} and PREDICTION_ROOT={PREDICTION_ROOT}")

    wb = Workbook()

    wb.remove(wb.active)

    summary_headers = (["dataset"]

                       + [f"IoU_{m}" for m in ("QL", "Ncut", "QPE", "IQPE", "Quantum")]

                       + [f"Dice_{m}" for m in ("QL", "Ncut", "QPE", "IQPE", "Quantum")]

                       + [f"Patch_count_{m}" for m in METHOD_NAMES]

                       + [f"Missing_GT_{m}" for m in METHOD_NAMES]

                       + [f"Missing_prediction_{m}" for m in METHOD_NAMES]

                       + [f"Patch_errors_{m}" for m in METHOD_NAMES])

    summary = make_sheet(wb, "Results", summary_headers)

    details = make_sheet(wb, "Patch details",

                         ["dataset", "method", "patch_id", "GT_file", "prediction_file",

                          "mIoU_%", "Macro_Dice_%", "status"])

    total_start = time.time()

    for dataset in dataset_names():

        gt_folder = GT_ROOT / dataset

        if not gt_folder.is_dir():

            print(f"Skipping {dataset}: ground-truth folder not found: {gt_folder}")

            continue

        try:

            gt_files = index_folder(gt_folder, dataset)

        except ValueError as exc:

            print(f"Skipping {dataset}: {exc}")

            continue

        result_row = {column: None for column in summary_headers}

        result_row["dataset"] = dataset

        for method in METHOD_NAMES:

            pred_folder = PREDICTION_ROOT / dataset / PREDICTION_SUBFOLDER / method

            if not pred_folder.is_dir():

                print(f"Skipping {dataset}/{method}: folder not found: {pred_folder}")

                continue

            try:

                pred_files = index_folder(pred_folder, dataset)

            except ValueError as exc:

                print(f"Skipping {dataset}/{method}: {exc}")

                continue

            gt_ids, pred_ids = set(gt_files), set(pred_files)

            miou_scores, dice_scores = [], []

            errors = 0

            for patch_id in sorted(gt_ids | pred_ids):

                gt_path, pred_path = gt_files.get(patch_id), pred_files.get(patch_id)

                if gt_path is None or pred_path is None:

                    status = "missing GT" if gt_path is None else "missing prediction"

                    details.append([dataset, method, patch_id,

                                    gt_path.name if gt_path else "",

                                    pred_path.name if pred_path else "", None, None, status])

                    continue

                try:

                    overlap = accumulate_overlap(gt_path, pred_path)

                    miou, dice, _ = calculate_macro_metrics(overlap, USE_MAPPING)

                    miou_scores.append(miou)

                    dice_scores.append(dice)

                    details.append([dataset, method, patch_id, gt_path.name,

                                    pred_path.name, miou, dice, "OK"])

                except Exception as exc:

                    errors += 1

                    details.append([dataset, method, patch_id, gt_path.name,

                                    pred_path.name, None, None, str(exc)])

                if PRINT_EVERY_PATCHES and len(miou_scores) % PRINT_EVERY_PATCHES == 0 and miou_scores:

                    print(f"{dataset}/{method}: {len(miou_scores)} patches completed")

            result_row[f"Patch_count_{method}"] = len(miou_scores)

            result_row[f"Missing_GT_{method}"] = len(pred_ids - gt_ids)

            result_row[f"Missing_prediction_{method}"] = len(gt_ids - pred_ids)

            result_row[f"Patch_errors_{method}"] = errors

            result_row[f"IoU_{method}"] = (sum(miou_scores) / len(miou_scores)

                                             if miou_scores else None)

            result_row[f"Dice_{method}"] = (sum(dice_scores) / len(dice_scores)

                                              if dice_scores else None)

            print(f"{dataset}/{method}: {len(miou_scores)} patch, "

                  f"mIoU={result_row[f'IoU_{method}']}, "

                  f"Macro-Dice={result_row[f'Dice_{method}']}")

        summary.append([result_row.get(col) for col in summary_headers])

    for sheet in (summary, details):

        sheet.auto_filter.ref = sheet.dimensions

        sheet.column_dimensions["A"].width = 53

        for i in range(2, sheet.max_column + 1):

            sheet.column_dimensions[get_column_letter(i)].width = 23 if i != 4 and i != 5 else 60

        for row in sheet.iter_rows(min_row=2):

            for cell in row:

                if isinstance(cell.value, float):

                    cell.number_format = "0.000000"

    OUTPUT_EXCEL.parent.mkdir(parents=True, exist_ok=True)

    wb.save(OUTPUT_EXCEL)

    print(f"Saved: {OUTPUT_EXCEL} (elapsed time {time.time()-total_start:.1f} seconds)")

if __name__ == "__main__":

    main()
